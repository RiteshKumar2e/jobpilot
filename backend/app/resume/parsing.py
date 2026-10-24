"""Read the candidate's real resume files and extract plain text + a variant label.

Supported: .pdf (pypdf), .docx (python-docx), .txt/.md, .tex (plain read).
The extracted text is the SOURCE OF TRUTH for what may appear on a tailored
resume — the tailoring step is not allowed to introduce skills/claims absent
from here (see tailor.py verify step).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List


@dataclass
class ParsedResume:
    path: str
    filename: str
    ext: str
    text: str
    variant_label: str = ""
    skills_found: List[str] = field(default_factory=list)


GENAI = "Generative AI / LLM"
AIML = "AI/ML"
ML = "Machine Learning"
CV = "Computer Vision / Research"
SD = "Software Development / Backend"
FULLSTACK = "Full-Stack / Python"
GENERAL = "General AI Engineer"

# Filename tokens are the most reliable signal: a resume named Resume_GENAI.pdf
# is the GenAI variant even though its text also mentions OpenCV/research.
# Order matters ("ai"+"ml" must win over plain "ml").
_FILENAME_RULES = [
    (GENAI, lambda t: bool(t & {"genai", "gen", "generative", "llm", "rag", "agentic"})),
    (AIML, lambda t: "aiml" in t or {"ai", "ml"} <= t),
    (FULLSTACK, lambda t: bool(t & {"fullstack", "full", "mern", "web", "frontend"})),
    (SD, lambda t: bool(t & {"sd", "sde", "swe", "software", "backend", "dev", "developer"})),
    (ML, lambda t: bool(t & {"ml", "machine", "datascience", "ds"})),
    (CV, lambda t: bool(t & {"vision", "research"})),
]

# Content fallback: the variant whose keywords occur MOST wins (first-match
# labelled every resume "CV/Research" because they all mention research).
_CONTENT_RULES = {
    GENAI: ["llm", "rag", "langchain", "generative", "genai", "prompt", "agent", "vector"],
    ML: ["machine learning", "scikit-learn", "regression", "classification", "xgboost", "pandas"],
    CV: ["computer vision", "opencv", "neu-det", "defect", "mobilenet", "segmentation", "publication"],
    SD: ["backend", "rest api", "microservice", "fastapi", "data structures", "system design"],
    FULLSTACK: ["react", "full stack", "full-stack", "frontend", "express", "next.js", "tailwind"],
}


def _label_from_filename(name: str) -> str:
    tokens = {t for t in re.split(r"[^a-z0-9]+", Path(name).stem.lower()) if t}
    tokens -= {"resume", "cv"}  # "CV" here usually means curriculum vitae
    return next((lbl for lbl, rule in _FILENAME_RULES if rule(tokens)), "")


def _label_from_content(low: str) -> str:
    counts = {lbl: sum(low.count(k) for k in kws) for lbl, kws in _CONTENT_RULES.items()}
    best = max(counts, key=counts.get)
    return best if counts[best] else GENERAL


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def parse_resume_file(path: str | Path, known_skills: List[str]) -> ParsedResume:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        text = _read_pdf(p)
    elif ext in (".docx",):
        text = _read_docx(p)
    elif ext in (".txt", ".md", ".tex"):
        text = _read_text(p)
    else:
        raise ValueError(f"Unsupported resume type: {ext}")

    low = text.lower()
    label = _label_from_filename(p.name) or _label_from_content(low)
    skills_found = sorted({s for s in known_skills if s.lower() in low})

    return ParsedResume(path=str(p), filename=p.name, ext=ext, text=text,
                        variant_label=label, skills_found=skills_found)


def discover_resumes(resumes_dir: str | Path, known_skills: List[str]) -> List[ParsedResume]:
    d = Path(resumes_dir)
    if not d.exists():
        return []
    out: List[ParsedResume] = []
    for p in sorted(d.iterdir()):
        if p.name.startswith(".") or p.is_dir():
            continue
        if p.suffix.lower() in (".pdf", ".docx", ".txt", ".md", ".tex"):
            try:
                out.append(parse_resume_file(p, known_skills))
            except Exception:
                continue
    return out
